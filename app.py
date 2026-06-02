"""
app.py — Flask backend para Deutsch Lernen
==========================================
Reemplaza el Streamlit app con Flask + API REST + Jinja2 templates.

Estructura de archivos:
    flask_app/
    ├── app.py                  ← este archivo
    ├── templates/
    │   ├── base.html
    │   ├── practica.html
    │   ├── completar.html
    │   ├── buscar.html
    │   └── verbos.html
    ├── static/
    │   └── style.css
    └── requirements.txt

Instalar:   pip install flask pandas openpyxl requests beautifulsoup4
Ejecutar:   python app.py
Acceder desde el celular (misma WiFi): http://<IP_LOCAL>:5000
"""

import os
import random
import shutil
import threading
from pathlib import Path
from urllib.parse import quote

import openpyxl
import pandas as pd
import requests as req
from bs4 import BeautifulSoup
from flask import Flask, jsonify, redirect, render_template, request, url_for

app = Flask(__name__)
app.config["JSON_AS_ASCII"] = False

# ── Ruta al Excel ──────────────────────────────────────────────
# Local:   usa ../data/dic.xlsx  (relativo a app.py)
# Railway: set DATA_FILE=/data/dic.xlsx  (volumen persistente)
_default_data = Path(__file__).parent.parent / "data" / "dic.xlsx"
DATA_FILE     = Path(os.environ.get("DATA_FILE", str(_default_data)))


def _ensure_data_file():
    """
    Garantiza que DATA_FILE existe antes de arrancar.

    En Railway (primera vez):
      - Copia dic.xlsx desde la versión bundled en el repo
        hacia el volumen persistente (/data/dic.xlsx).
    En despliegues posteriores:
      - El archivo YA existe en el volumen → no hace nada.
      - Los datos del usuario se conservan entre deploys.
    Localmente:
      - El archivo ya está en ../data/dic.xlsx → no hace nada.
    """
    if DATA_FILE.exists():
        return  # ya existe, no sobreescribir

    # Buscar versión inicial en orden de prioridad
    bundled_candidates = [
        Path(__file__).parent / "dic.xlsx",                    # junto a app.py (Railway)
        Path(__file__).parent.parent / "data" / "dic.xlsx",   # ruta local dev
    ]
    for src in bundled_candidates:
        if src.exists():
            DATA_FILE.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(src, DATA_FILE)
            print(f"📋 dic.xlsx copiado: {src} → {DATA_FILE}")
            return

    raise FileNotFoundError(
        f"No se encontró dic.xlsx. "
        f"Colócalo junto a app.py o define DATA_FILE. "
        f"Rutas buscadas: {[str(c) for c in bundled_candidates]}"
    )


_ensure_data_file()   # ejecutar al importar el módulo
SHEET_WORTE  = "Wörte"
SHEET_VERBEN = "Verben"
FILL_COLS    = [
    "Artikel", "Plural nominativ", "Singular genitiv",
    "Spanisch 1", "Spanisch 2", "Spanisch 3",
]

# ── Caché ─────────────────────────────────────────────────────────
_df_cache   = {}
_df_lock    = threading.Lock()
_conj_cache = {}


def _load(sheet: str) -> pd.DataFrame:
    with _df_lock:
        if sheet not in _df_cache:
            df = pd.read_excel(DATA_FILE, sheet_name=sheet, dtype=str)
            df.columns = [c.strip() for c in df.columns]
            df["Orden"] = pd.to_numeric(df["Orden"], errors="coerce")
            df = df.dropna(subset=["Orden"])
            df["Orden"] = df["Orden"].astype(int)
            df = df.fillna("")
            for col in df.columns:
                if df[col].dtype == object:
                    df[col] = df[col].str.strip()
            _df_cache[sheet] = df.sort_values("Orden").reset_index(drop=True)
    return _df_cache[sheet]


def _invalidate(sheet: str):
    with _df_lock:
        _df_cache.pop(sheet, None)


# ── Escritura Excel ───────────────────────────────────────────────
def _update_cell(sheet, orden, column, value):
    try:
        wb = openpyxl.load_workbook(DATA_FILE)
        ws = wb[sheet]
        hdrs = {str(c.value).strip(): c.column for c in ws[1] if c.value}
        if column not in hdrs:
            return False, f"Columna '{column}' no encontrada."
        tcol = hdrs[column]
        ocol = hdrs.get("Orden", 1)
        for row in ws.iter_rows(min_row=2):
            raw = row[ocol - 1].value
            if raw is None:
                continue
            try:
                if int(float(str(raw))) == orden:
                    ws.cell(row=row[0].row, column=tcol).value = value or None
                    wb.save(DATA_FILE)
                    _invalidate(sheet)
                    return True, "Guardado correctamente."
            except (ValueError, TypeError):
                continue
        return False, f"Orden {orden} no encontrado."
    except Exception as exc:
        return False, str(exc)


def _append_row(sheet, fields):
    wb = openpyxl.load_workbook(DATA_FILE)
    ws = wb[sheet]
    hdrs = {str(c.value).strip(): c.column for c in ws[1] if c.value}
    ocol = hdrs.get("Orden", 1)
    max_ord = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        val = row[ocol - 1]
        if val:
            try:
                max_ord = max(max_ord, int(float(str(val))))
            except (ValueError, TypeError):
                pass
    new_ord = max_ord + 1
    new_row = [None] * ws.max_column
    for col_name, col_idx in hdrs.items():
        if col_name == "Orden":
            new_row[col_idx - 1] = new_ord
        elif col_name in fields:
            new_row[col_idx - 1] = fields[col_name] or None
    ws.append(new_row)
    wb.save(DATA_FILE)
    _invalidate(sheet)
    return new_ord


# ── Reverso helpers ───────────────────────────────────────────────
_REV_H    = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)",
             "Accept-Language": "de-DE,de;q=0.9"}
_SKIP_M   = {"Infinitiv"}
_MOODS    = ["Konjunktiv II", "Konjunktiv I", "Indikativ", "Imperativ", "Partizip"]


def _reverso_url(verb):
    return f"https://conjugator.reverso.net/conjugation-german-verb-{quote(verb.lower(), safe='')}.html"


def _mood_tense(title):
    for m in _MOODS:
        if title.startswith(m):
            t = title[len(m):].strip()
            return m, (t or m)
    return "Andere", title


def _parse_li(li):
    pron, segs = "", []
    for tag in li.find_all("i"):
        cls  = tag.get("class", [])
        text = tag.get_text(strip=True)
        if not text:
            continue
        if "graytxt" in cls:
            pron = text.strip("(). ")
        elif "particletxt" in cls:
            segs.append(("p", text))
        elif "verbtxt" in cls:
            if segs and segs[-1][0] == "p":
                _, p = segs.pop()
                segs.append(("v", p + text))
            else:
                segs.append(("v", text))
        elif "auxgraytxt" in cls:
            segs.append(("a", text))
    return pron, " ".join(t for _, t in segs).strip()


def _fetch_conj(verb, url=""):
    key = url or verb
    if key in _conj_cache:
        return _conj_cache[key]
    target = url or _reverso_url(verb)
    try:
        resp = req.get(target, headers=_REV_H, timeout=14)
        resp.raise_for_status()
    except Exception as exc:
        return {"_error": str(exc)}
    soup   = BeautifulSoup(resp.text, "html.parser")
    result = {}
    for block in soup.find_all("div", class_="blue-box-wrap"):
        mt = block.get("mobile-title", "").strip()
        if not mt or any(mt.startswith(s) for s in _SKIP_M):
            continue
        mood, tense = _mood_tense(mt)
        ul = block.find("ul", class_="wrap-verbs-listing")
        if not ul:
            continue
        pairs = []
        for li in ul.find_all("li"):
            p, f = _parse_li(li)
            if f:
                pairs.append({"pron": p, "form": f})
        if pairs:
            result.setdefault(mood, {})[tense] = pairs
    _conj_cache[key] = result
    return result


# ══════════════════════════════════════════════════════════════════
# PÁGINAS
# ══════════════════════════════════════════════════════════════════
@app.route("/")
def index():
    return redirect(url_for("practica"))


@app.route("/practica")
def practica():
    df      = _load(SHEET_WORTE)
    min_ord = int(df["Orden"].min())
    max_ord = int(df["Orden"].max())
    orden   = max(min_ord, min(int(request.args.get("orden", min_ord)), max_ord))
    return render_template("practica.html",
                           orden=orden, min_orden=min_ord, max_orden=max_ord)


@app.route("/completar")
def completar():
    return render_template("completar.html", fill_cols=FILL_COLS)


@app.route("/buscar")
def buscar():
    return render_template("buscar.html")


@app.route("/verbos")
def verbos():
    df = _load(SHEET_VERBEN)
    verb_list = [
        {"orden": int(r["Orden"]), "infinitiv": r["Infinitiv"],
         "bedeutung": r.get("Bedeutung", ""), "link": r.get("Link", "")}
        for _, r in df.iterrows() if r["Infinitiv"]
    ]
    return render_template("verbos.html", verb_list=verb_list)


# ══════════════════════════════════════════════════════════════════
# API — WÖRTE
# ══════════════════════════════════════════════════════════════════
@app.route("/api/word/<int:orden>")
def api_word(orden):
    df     = _load(SHEET_WORTE)
    row_df = df[df["Orden"] == orden]
    if row_df.empty:
        return jsonify({"error": "Not found"}), 404
    r = row_df.iloc[0]
    return jsonify({
        "orden":    int(r["Orden"]),
        "singular": r.get("Singular nominativ", ""),
        "artikel":  r.get("Artikel", ""),
        "plural":   r.get("Plural nominativ", ""),
        "genitiv":  r.get("Singular genitiv", ""),
        "span1":    r.get("Spanisch 1", ""),
        "span2":    r.get("Spanisch 2", ""),
        "span3":    r.get("Spanisch 3", ""),
    })


@app.route("/api/word/random")
def api_word_random():
    df = _load(SHEET_WORTE)
    return jsonify({"orden": int(random.choice(df["Orden"].tolist()))})


@app.route("/api/word/save", methods=["POST"])
def api_word_save():
    d       = request.json
    ok, msg = _update_cell(SHEET_WORTE, int(d["orden"]), d["column"], d.get("value", ""))
    return jsonify({"ok": ok, "msg": msg})


@app.route("/api/word/add", methods=["POST"])
def api_word_add():
    try:
        new_ord = _append_row(SHEET_WORTE, request.json)
        return jsonify({"ok": True, "orden": new_ord})
    except Exception as exc:
        return jsonify({"ok": False, "msg": str(exc)})


@app.route("/api/empty_cells")
def api_empty_cells():
    df     = _load(SHEET_WORTE)
    result = []
    for _, row in df.iterrows():
        for col in FILL_COLS:
            if row.get(col, "") == "":
                result.append({
                    "orden":    int(row["Orden"]),
                    "singular": row.get("Singular nominativ", ""),
                    "artikel":  row.get("Artikel", ""),
                    "columna":  col,
                })
    return jsonify(result)


@app.route("/api/search/es")
def api_search_es():
    q  = request.args.get("q", "").strip().lower()
    df = _load(SHEET_WORTE)
    if not q:
        return jsonify([])
    mask = (df["Spanisch 1"].str.lower().str.contains(q, na=False) |
            df["Spanisch 2"].str.lower().str.contains(q, na=False) |
            df["Spanisch 3"].str.lower().str.contains(q, na=False))
    out = []
    for _, r in df[mask].iterrows():
        sv = [r.get(f"Spanisch {i}", "") for i in [1, 2, 3]]
        out.append({
            "orden":   int(r["Orden"]), "artikel": r.get("Artikel", ""),
            "singular":r.get("Singular nominativ", ""),
            "plural":  r.get("Plural nominativ", ""),
            "genitiv": r.get("Singular genitiv", ""),
            "span":    " ; ".join(s for s in sv if s),
        })
    return jsonify(out)


@app.route("/api/search/de")
def api_search_de():
    q  = request.args.get("q", "").strip().lower()
    df = _load(SHEET_WORTE)
    if not q:
        return jsonify([])
    mask = (df["Singular nominativ"].str.lower().str.contains(q, na=False) |
            df["Plural nominativ"].str.lower().str.contains(q, na=False))
    out = []
    for _, r in df[mask].iterrows():
        sv = [r.get(f"Spanisch {i}", "") for i in [1, 2, 3]]
        out.append({
            "orden":        int(r["Orden"]), "artikel": r.get("Artikel", ""),
            "singular":     r.get("Singular nominativ", ""),
            "plural":       r.get("Plural nominativ", ""),
            "genitiv":      r.get("Singular genitiv", ""),
            "translations": [s for s in sv if s],
        })
    return jsonify(out)


# ══════════════════════════════════════════════════════════════════
# API — VERBEN
# ══════════════════════════════════════════════════════════════════
@app.route("/api/verb/conjugations")
def api_verb_conjugations():
    verb = request.args.get("verb", "").strip()
    url  = request.args.get("url",  "").strip()
    if not verb:
        return jsonify({"_error": "verb param required"}), 400
    return jsonify(_fetch_conj(verb, url))


@app.route("/api/verb/save", methods=["POST"])
def api_verb_save():
    d      = request.json
    orden  = int(d.get("orden", 0))
    errors = []
    for col in ["Infinitiv", "Bedeutung", "Link"]:
        val = d.get(col)
        if val is not None:
            ok, msg = _update_cell(SHEET_VERBEN, orden, col, val)
            if not ok:
                errors.append(msg)
    _conj_cache.pop(d.get("link", ""), None)
    return jsonify({"ok": not errors, "errors": errors})


# ══════════════════════════════════════════════════════════════════
if __name__ == "__main__":
    print("\n🇩🇪  Deutsch Lernen — Flask")
    print(f"   Excel:  {DATA_FILE}")
    print("   Local:  http://127.0.0.1:5000")
    print("   Celular (misma WiFi): http://<TU_IP>:5000\n")
    _port = int(os.environ.get("PORT", 5000))
    _debug = os.environ.get("FLASK_ENV") != "production"
    app.run(host="0.0.0.0", port=_port, debug=_debug)
